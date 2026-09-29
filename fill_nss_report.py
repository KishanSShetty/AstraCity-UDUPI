#!/usr/bin/env python3
"""
Fill the NSS one-page report — keep original formatting, expand content slightly
but stay within one page.
"""

from docx import Document
import os

src = r"C:\Users\Kishan Shetty\Downloads\AstraSky-maing\NSS_1 (1).docx"
doc = Document(src)

def set_text(p, text):
    if not p.runs:
        p.text = text
        return
    first = p.runs[0]
    bold = first.bold
    size = first.font.size
    name = first.font.name
    color = first.font.color.rgb if first.font.color and first.font.color.rgb else None
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = size
    r.font.name = name
    if color:
        r.font.color.rgb = color

paras = doc.paragraphs

# P2: Date
set_text(paras[2], "Date: January 2026 - June 2026 (Even Semester)")

# P3: NSS intro — keep original template text as-is

# P4: Objectives
set_text(paras[4],
    "Objectives: "
    "(1) To conduct a comprehensive, data-driven audit of Solid Waste Management (SWM) in HSR Layout, "
    "Ward 174, Bengaluru using space technology, GIS mapping, and AI-based analytics. "
    "(2) To promote environmental awareness among students and the community through Environment Day "
    "celebrations, tree plantation, and sustainability project demonstrations."
)

# P5: Activities Undertaken
set_text(paras[5],
    "Activities Undertaken: "
    "Activity 1 - SWM Audit using AstraCity Digital Twin Platform: "
    "Conducted a full-scale audit of solid waste management in HSR Layout (Ward 174, Bengaluru). "
    "The audit used Sentinel-2 satellite imagery, QGIS for land-use classification, OpenStreetMap data "
    "for road network and building analysis, and AI algorithms including Clarke-Wright CVRP for route "
    "optimisation, IPCC Methane Model for emissions estimation, and DBSCAN for illegal dump detection. "
    "Key findings: HSR generates 55 T/day of waste with 0 dumpyards, 0 BMUs, and only 6 DWCCs serving "
    "1,10,000 residents. Route optimisation reduced daily distance from 132.44 km to 32.41 km (75.5% "
    "improvement). Four illegal dump sites totalling 793 sq.m were detected. Estimated annual savings: "
    "Rs.9.42 Crores for this single ward. "
    "Activity 2 - World Environment Day Celebration: "
    "Participated in the Environment Day event organised by NSS, RVCE. Showcased the AstraCity digital "
    "twin platform to fellow students and faculty, demonstrating how technology can address urban waste "
    "challenges. Brought and distributed plants among students. Other students also brought saplings "
    "and participated in a campus tree plantation drive."
)

# P6: Activities Participated
set_text(paras[6],
    "Activities Participated: "
    "(1) SWM Field Audit and Data Collection in HSR Layout, Ward 174 - GIS mapping, satellite image "
    "analysis, road and building surveys, and waste facility cataloguing. "
    "(2) World Environment Day - Project exhibition of AstraCity platform, plant distribution, "
    "and campus tree plantation drive."
)

# P7: Community Service
set_text(paras[7],
    "Community Service: "
    "The SWM audit directly serves the community of HSR Layout (~1,10,000 residents) by identifying "
    "critical infrastructure deficits - no dumpyards, no BMUs, and overloaded DWCCs. The audit report "
    "includes actionable recommendations such as optimised collection routes, illegal dump cleanup "
    "priorities, and bio-methanisation potential at Kudlu. The report has been prepared for submission "
    "to BBMP. The Environment Day plantation activity contributed to campus greening and promoted "
    "sustainability awareness among students."
)

# P8: Ability Enhancement / Learnings
set_text(paras[8],
    "Ability Enhancement/Learnings: "
    "(1) Gained hands-on experience with GIS tools (QGIS, Sentinel-2 satellite data, OpenStreetMap). "
    "(2) Applied AI/ML algorithms (CVRP routing, DBSCAN clustering, IPCC methane modelling) to "
    "real-world civic problems. "
    "(3) Developed data-driven policy analysis skills - translating technical findings into actionable "
    "recommendations for municipal authorities. "
    "(4) Strengthened teamwork, public presentation, and environmental awareness through the "
    "Environment Day showcase and plantation activity."
)

# P9-P13: Personal details
set_text(paras[9],  "Name: Kishan Shetty")
set_text(paras[10], "USN: 1RV24IS058")
set_text(paras[11], "Dept: Information Science and Engineering (ISE)")
set_text(paras[12], "E mail: kishans.is24@rvce.edu.in")
set_text(paras[13], "Phone Number: ")

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "NSS_Report_Filled.docx")
doc.save(out)
print(f"[OK] Report saved to: {out}")
