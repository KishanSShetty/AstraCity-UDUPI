import docx

doc = docx.Document()
with open('AstraCity_Formatted_Paper.md', 'r', encoding='utf-8') as f:
    lines = f.readlines()

in_table = False
table_data = []

for line in lines:
    line = line.strip()
    
    if line.startswith('|'):
        if '---' in line:
            continue # separator
        row = [cell.strip() for cell in line.split('|')[1:-1]]
        table_data.append(row)
        in_table = True
        continue
    else:
        if in_table:
            # create table
            table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
            table.style = 'Table Grid'
            for i, row in enumerate(table_data):
                for j, cell in enumerate(row):
                    table.cell(i, j).text = cell.replace('**', '')
            table_data = []
            in_table = False

    if not line:
        continue

    if line.startswith('# '):
        doc.add_heading(line[2:], 0)
    elif line.startswith('## '):
        doc.add_heading(line[3:], 1)
    elif line.startswith('### '):
        doc.add_heading(line[4:], 2)
    elif line.startswith('- '):
        doc.add_paragraph(line[2:].replace('**', ''), style='List Bullet')
    elif line == '---':
        pass
    else:
        p = doc.add_paragraph()
        parts = line.split('**')
        for i, part in enumerate(parts):
            run = p.add_run(part.replace('`', ''))
            if i % 2 == 1:
                run.bold = True

if in_table:
    table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
    table.style = 'Table Grid'
    for i, row in enumerate(table_data):
        for j, cell in enumerate(row):
            table.cell(i, j).text = cell.replace('**', '')

doc.save('AstraCity_Formatted_Paper.docx')
