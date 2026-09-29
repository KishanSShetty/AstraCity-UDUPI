import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

doc = Document(r'C:\Users\Kishan Shetty\Downloads\AstraSky-maing\AstraCityssss_Conference_Final.docx')

print("=== PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    style = p.style.name if p.style else 'None'
    text = p.text[:100] if p.text else ''
    if text.strip():
        print(f'[{i}] Style=[{style}] Text=[{text}]')

print(f"\n=== TABLES: {len(doc.tables)} ===")
for ti, tbl in enumerate(doc.tables):
    print(f'  Table {ti}: {len(tbl.rows)} rows x {len(tbl.columns)} cols')
    for ri, row in enumerate(tbl.rows):
        cells = [c.text[:30] for c in row.cells]
        print(f'    Row {ri}: {cells}')
