import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn

doc = Document(r'C:\Users\Kishan Shetty\Downloads\AstraSky-maing\template_clean.docx')

# 1. List all styles used
print('=== STYLES USED IN TEMPLATE ===')
for p in doc.paragraphs:
    if p.text.strip():
        style_name = p.style.name
        txt = p.text[:90].replace('\n', ' ')
        print(f'  Style: [{style_name}]  Text: [{txt}]')

# 2. Section properties
print('\n=== SECTIONS ===')
for i, section in enumerate(doc.sections):
    print(f'Section {i}: width={section.page_width}, height={section.page_height}')
    print(f'  margins: top={section.top_margin}, bottom={section.bottom_margin}, left={section.left_margin}, right={section.right_margin}')
    sectPr = section._sectPr
    cols = sectPr.find(qn('w:cols'))
    if cols is not None:
        num = cols.get(qn('w:num'))
        space = cols.get(qn('w:space'))
        eq = cols.get(qn('w:equalWidth'))
        print(f'  columns: num={num}, space={space}, equalWidth={eq}')
    else:
        print('  columns: NONE (single column)')

# 3. Tables
print(f'\n=== TABLES: {len(doc.tables)} ===')
for i, t in enumerate(doc.tables):
    print(f'  Table {i}: {len(t.rows)} rows x {len(t.columns)} cols')
    for j, row in enumerate(t.rows):
        cells_text = [c.text[:40] for c in row.cells]
        print(f'    Row {j}: {cells_text}')

# 4. All available styles in template
print('\n=== ALL DEFINED STYLES ===')
for s in doc.styles:
    if s.type is not None:
        print(f'  {s.style_id} -> {s.name} (type={s.type})')
