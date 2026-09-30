import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn

doc = Document(r'C:\Users\Kishan Shetty\Downloads\AstraSky-maing\AstraCity_Conference_Final.docx')

print('=== SECTIONS ===')
for i, section in enumerate(doc.sections):
    sectPr = section._sectPr
    cols = sectPr.find(qn('w:cols'))
    if cols is not None:
        num = cols.get(qn('w:num'))
        space = cols.get(qn('w:space'))
        print(f'Section {i}: columns num={num}, space={space}')
    else:
        print(f'Section {i}: single column')

print(f'\n=== PARAGRAPHS: {len(doc.paragraphs)} ===')
for idx, p in enumerate(doc.paragraphs):
    if p.text.strip():
        txt = p.text[:80].replace('\n', ' ')
        print(f'[{idx}] Style=[{p.style.name}] Align=[{p.alignment}] Text=[{txt}]')

print(f'\n=== TABLES: {len(doc.tables)} ===')
for i, t in enumerate(doc.tables):
    print(f'  Table {i}: {len(t.rows)} rows x {len(t.columns)} cols')
    for j, row in enumerate(t.rows):
        cells = [c.text[:30] for c in row.cells]
        print(f'    Row {j}: {cells}')
